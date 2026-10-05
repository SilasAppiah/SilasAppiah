import json, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment

links=json.load(open('whatsapp-links.json'))
NAVY="171365"; PALE="EEF2FA"; INPUT="FFF2CC"; MUTED="5B6180"
F=lambda **k: Font(name="Arial", size=k.pop("size",10), **k)
HDR_FILL=PatternFill("solid", fgColor=NAVY); IN_FILL=PatternFill("solid", fgColor=INPUT); PALE_FILL=PatternFill("solid", fgColor=PALE)
thin=Side(style="thin", color="D6E0F2"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True, vertical="top")

wb=Workbook()
def title(ws, t, sub):
    ws["A1"]=t; ws["A1"].font=F(size=14,bold=True,color=NAVY)
    ws["A2"]=sub; ws["A2"].font=F(size=10,italic=True,color=MUTED)
def header(ws,row,cols):
    for i,c in enumerate(cols,1):
        cell=ws.cell(row=row,column=i,value=c); cell.font=F(bold=True,color="FFFFFF"); cell.fill=HDR_FILL
        cell.alignment=Alignment(wrap_text=True,vertical="center"); cell.border=BOX
def style(cell, inp=False, fmt=None, bold=False):
    cell.font=F(color="0000FF" if inp else "000000", bold=bold); cell.border=BOX
    if inp: cell.fill=IN_FILL
    if fmt: cell.number_format=fmt

# ---------- Read me ----------
ws=wb.active; ws.title="Read me"
title(ws,"Social Enquiry Routing: tracker","Innovation 2, Q4 2026. Owner: Silas Atuahene Appiah, Digital Marketing Officer")
rows=[("Sheet","What it's for","When"),
("WhatsApp links","All 15 tagged links to Sales (0244 122 855). Copy into saved replies, automations, bios and posts.","Set up once"),
("UTM builder","Builds tracked links to the website or portal so Google Analytics shows the source.","Every time you post a website link"),
("Baseline DMs","Reply times for 20 recent DMs, taken before the automations go live.","Week of 5 October, again in December"),
("Weekly log","Enquiries, automation, tagged handovers and enrolments, one row per week.","Every Friday, 10 minutes"),
("Dashboard","Q4 results against targets. Updates itself from the other sheets.","Monthly report")]
for r,row in enumerate(rows,4):
    for c,v in enumerate(row,1):
        cell=ws.cell(row=r,column=c,value=v); cell.alignment=wrap; cell.border=BOX
        cell.font=F(bold=(r==4),color="FFFFFF" if r==4 else "000000")
        if r==4: cell.fill=HDR_FILL
ws["A12"]="Legend"; ws["A12"].font=F(bold=True,color=NAVY)
ws["A13"]="Yellow cell, blue text"; ws["A13"].fill=IN_FILL; ws["A13"].font=F(color="0000FF"); ws["B13"]="You fill this in"
ws["A14"]="White cell, black text"; ws["A14"].font=F(); ws["B14"]="Formula, leave it alone"
ws["A15"]="Grey italic row"; ws["A15"].font=F(italic=True,color=MUTED); ws["B15"]="Example showing the format. Not counted."
ws["A17"]="Ref code key"; ws["A17"].font=F(bold=True,color=NAVY)
ws["A18"]="FB / IG / LI"; ws["B18"]="Facebook / Instagram / LinkedIn"
ws["A19"]="APHRI / PHRI / SPHRI"; ws["B19"]="Programme the person asked about"
ws["A20"]="CONSULT / GEN"; ws["B20"]="HRCC Consulting / general enquiry"
for c in ("A18","A19","A20","B13","B14","B15","B18","B19","B20"): ws[c].font=F()
ws.column_dimensions["A"].width=24; ws.column_dimensions["B"].width=70; ws.column_dimensions["C"].width=30

# ---------- WhatsApp links ----------
ws=wb.create_sheet("WhatsApp links")
title(ws,"Tagged WhatsApp links to Sales","Number: 0244 122 855. Each link opens WhatsApp with the message and ref code already typed. Test every link once before use.")
header(ws,4,["Platform","Programme","Ref code","Message the person sends","Link"])
for i,l in enumerate(links,5):
    vals=[l["platform"],l["programme"],l["ref"],l["message"],l["link"]]
    for c,v in enumerate(vals,1):
        cell=ws.cell(row=i,column=c,value=v); style(cell); cell.alignment=Alignment(vertical="top",wrap_text=(c==4))
    ws.cell(row=i,column=5).hyperlink=l["link"]; ws.cell(row=i,column=5).font=F(color="0563C1",underline="single")
for col,w in zip("ABCDE",[12,17,13,52,60]): ws.column_dimensions[col].width=w
ws.freeze_panes="A5"

# ---------- UTM builder ----------
ws=wb.create_sheet("UTM builder")
title(ws,"UTM link builder","Fill the yellow cells. Column F builds the link. Keep everything lower case so Google Analytics groups it correctly.")
header(ws,4,["Base URL (page on the website or portal)","Source","Medium","Campaign","Content (programme_posttype)","Tracked link (copy this)"])
ex=["https://www.phrglobal.com/","instagram","social","cohort49","phri_classroom"]
for c,v in enumerate(ex,1):
    cell=ws.cell(row=5,column=c,value=v); cell.font=F(italic=True,color=MUTED); cell.border=BOX
dvs=DataValidation(type="list",formula1='"facebook,instagram,linkedin"',allow_blank=True)
dvm=DataValidation(type="list",formula1='"social,paid_social"',allow_blank=True)
ws.add_data_validation(dvs); ws.add_data_validation(dvm)
def utm(r): return (f'=IF(A{r}="","",A{r}&IF(ISNUMBER(FIND("?",A{r})),"&","?")&"utm_source="&LOWER(B{r})&"&utm_medium="&LOWER(C{r})'
                    f'&"&utm_campaign="&LOWER(D{r})&IF(E{r}="","","&utm_content="&LOWER(E{r})))')
ws.cell(row=5,column=6,value=utm(5)).font=F(italic=True,color=MUTED); ws.cell(row=5,column=6).border=BOX
ws["G5"]="Example row"; ws["G5"].font=F(italic=True,color=MUTED)
for r in range(6,31):
    for c in range(1,6):
        cell=ws.cell(row=r,column=c); style(cell,inp=True)
        if c==4: cell.value="cohort49"
    style(ws.cell(row=r,column=6,value=utm(r)))
    dvs.add(f"B{r}"); dvm.add(f"C{r}")
for col,w in zip("ABCDEF",[38,12,13,12,24,90]): ws.column_dimensions[col].width=w
ws.freeze_panes="A6"

# ---------- Baseline DMs ----------
ws=wb.create_sheet("Baseline DMs")
title(ws,"Baseline: reply time for 20 recent DMs","Take the last 20 DMs that asked about a programme. Enter when each arrived and when someone first replied personally (not an automated reply).")
header(ws,4,["#","Platform","Received (date and time)","First personal reply (date and time)","Minutes to reply","Within 15 minutes?","Question asked"])
FMT="dd mmm yyyy hh:mm"
ex=["Ex.","Instagram",dt.datetime(2026,9,28,10,5),dt.datetime(2026,9,28,11,40),None,None,"PHRi fees"]
for c,v in enumerate(ex,1):
    cell=ws.cell(row=5,column=c,value=v); cell.font=F(italic=True,color=MUTED); cell.border=BOX
    if c in (3,4): cell.number_format=FMT
ws["E5"]='=IF(AND(C5<>"",D5<>""),ROUND((D5-C5)*1440,0),"")'; ws["F5"]='=IF(E5="","",IF(E5<=15,"Yes","No"))'
for c in ("E5","F5"): ws[c].font=F(italic=True,color=MUTED); ws[c].border=BOX
dvp=DataValidation(type="list",formula1='"Facebook,Instagram"',allow_blank=True); ws.add_data_validation(dvp)
for r in range(6,26):
    style(ws.cell(row=r,column=1,value=r-5))
    for c in (2,3,4,7): style(ws.cell(row=r,column=c),inp=True,fmt=FMT if c in (3,4) else None)
    dvp.add(f"B{r}")
    style(ws.cell(row=r,column=5,value=f'=IF(AND(C{r}<>"",D{r}<>""),ROUND((D{r}-C{r})*1440,0),"")'),fmt="0")
    style(ws.cell(row=r,column=6,value=f'=IF(E{r}="","",IF(E{r}<=15,"Yes","No"))'))
summ=[("DMs entered","=COUNT(E6:E25)","0"),("Average minutes to reply",'=IFERROR(AVERAGE(E6:E25),"")',"0"),
      ("Median minutes to reply",'=IFERROR(MEDIAN(E6:E25),"")',"0"),("Share replied within 15 minutes",'=IFERROR(COUNTIF(F6:F25,"Yes")/COUNT(E6:E25),"")',"0%")]
for i,(lab,f,fmt) in enumerate(summ,27):
    ws.cell(row=i,column=4,value=lab).font=F(bold=True,color=NAVY)
    style(ws.cell(row=i,column=5,value=f),fmt=fmt,bold=True)
ws["A32"]="Note: minutes are clock time, so a DM sent at night counts the hours until morning. That is fine for a baseline, as long as the December re-check is measured the same way."
ws["A32"].font=F(italic=True,color=MUTED,size=9)
for col,w in zip("ABCDEFG",[6,12,22,26,15,15,30]): ws.column_dimensions[col].width=w
ws.freeze_panes="A6"

# ---------- Weekly log ----------
ws=wb.create_sheet("Weekly log")
title(ws,"Weekly log","Fill one row every Friday. Tagged chat and enrolment numbers come from Sales; enquiries and automation counts from Meta Business Suite.")
cols=["Week starting (Monday)","Enquiries (DMs + comments)","Fully answered by automation","Share automated","Tagged chats: FB","Tagged chats: IG","Tagged chats: LI","Tagged chats: total","Social leads enrolled","Keyword-triggered DMs","Hours on repeat replies","Notes"]
header(ws,4,cols); ws.row_dimensions[4].height=42
start=dt.date(2026,10,5)
for i in range(13):
    r=5+i
    style(ws.cell(row=r,column=1,value=start+dt.timedelta(weeks=i)),fmt="dd mmm yyyy")
    for c in (2,3,5,6,7,9,10,11,12): style(ws.cell(row=r,column=c),inp=True,fmt="0.0" if c==11 else ("0" if c<12 else None))
    style(ws.cell(row=r,column=4,value=f'=IF(OR(B{r}="",B{r}=0,C{r}=""),"",C{r}/B{r})'),fmt="0%")
    style(ws.cell(row=r,column=8,value=f'=IF(COUNT(E{r}:G{r})=0,"",SUM(E{r}:G{r}))'),fmt="0")
T=18
ws.cell(row=T,column=1,value="Q4 total").font=F(bold=True,color=NAVY)
for c,L in zip((2,3,5,6,7,8,9,10),"BCEFGHIJ"):
    style(ws.cell(row=T,column=c,value=f"=SUM({L}5:{L}17)"),fmt="0",bold=True)
style(ws.cell(row=T,column=4,value='=IF(B18=0,"",C18/B18)'),fmt="0%",bold=True)
style(ws.cell(row=T,column=11,value='=IFERROR(AVERAGE(K5:K17),"")'),fmt="0.0",bold=True)
ws.cell(row=19,column=11,value="(weekly average)").font=F(italic=True,color=MUTED,size=9)
for col,w in zip("ABCDEFGHIJKL",[16,13,13,11,10,10,10,11,11,12,12,30]): ws.column_dimensions[col].width=w
ws.freeze_panes="B5"

# ---------- Dashboard ----------
ws=wb.create_sheet("Dashboard")
title(ws,"Q4 dashboard: Social Enquiry Routing","Updates from the other sheets. Only the yellow December cells need typing.")
header(ws,4,["Measure","Baseline","Q4 so far","Q4 target","On target?"])
WL="'Weekly log'"; BL="'Baseline DMs'"
dash=[
("Average minutes to a personal reply",f"={BL}!E28",None,15,"0","le","Re-sample 20 DMs in December and type the average here"),
("Share of DMs replied within 15 minutes",f"={BL}!E30",None,0.8,"0%","ge","Re-sample in December and type the share here"),
("Share of enquiries answered by automation",0,f"={WL}!D18",0.6,"0%","ge",None),
("Tagged handovers to Sales",0,f"={WL}!H18",None,"0",None,None),
("Social leads enrolled",0,f"={WL}!I18",None,"0",None,None),
("Handover to enrolment rate",None,f'=IFERROR({WL}!I18/{WL}!H18,"")',None,"0%",None,None),
("Hours a week on repeat replies",f'=IF({WL}!K5="","",{WL}!K5)',f'=IFERROR(INDEX({WL}!K5:K17,MATCH(9.99E+307,{WL}!K5:K17)),"")','=IF(B{r}="","",B{r}*2/3)',"0.0","le",None),
("Keyword-triggered DMs",0,f"={WL}!J18",None,"0",None,None),
]
for i,(lab,base,cur,tgt,fmt,cmp,note) in enumerate(dash,5):
    r=i+2 if False else i
    ws.cell(row=r,column=1,value=lab).font=F(); ws.cell(row=r,column=1).border=BOX
    style(ws.cell(row=r,column=2,value=base),fmt=fmt)
    cc=ws.cell(row=r,column=3,value=cur)
    if cur is None: style(cc,inp=True,fmt=fmt); cc.comment=Comment(note,"Claude")
    else: style(cc,fmt=fmt)
    if isinstance(tgt,str): tgt=tgt.format(r=r)
    style(ws.cell(row=r,column=4,value=tgt if tgt is not None else "Track"),fmt=fmt)
    if cmp:
        op="<=" if cmp=="le" else ">="
        style(ws.cell(row=r,column=5,value=f'=IF(OR(C{r}="",D{r}=""),"",IF(C{r}{op}D{r},"Yes","Not yet"))'))
    else: style(ws.cell(row=r,column=5,value=""))
ws["A14"]="Tagged handovers by platform"; ws["A14"].font=F(bold=True,color=NAVY)
for j,(p,L) in enumerate([("Facebook","E"),("Instagram","F"),("LinkedIn","G")],15):
    ws.cell(row=j,column=1,value=p).font=F(); ws.cell(row=j,column=1).border=BOX
    style(ws.cell(row=j,column=2,value=f"={WL}!{L}18"),fmt="0")
    style(ws.cell(row=j,column=3,value=f'=IFERROR(B{j}/SUM($B$15:$B$17),"")'),fmt="0%")
ws["B14"]="Chats"; ws["C14"]="Share"
for c in ("B14","C14"): ws[c].font=F(bold=True,color=NAVY)
ws["A19"]="Targets come from the Q4 innovation brief. The 80% within-15-minutes target is a working figure; adjust once the baseline is in."
ws["A19"].font=F(italic=True,color=MUTED,size=9)
for col,w in zip("ABCDE",[42,14,14,14,13]): ws.column_dimensions[col].width=w
for s in wb.worksheets: s.sheet_view.showGridLines=False
wb.save("Enquiry-Tracker.xlsx"); print("saved")
